"""Load ParcelPilot_Assessment_Data.xlsx into the SQLite structured-data layer.

Reads the workbook, validates its shape (sheets, columns, timestamp formats,
account references), and loads it into data/processed/parcelpilot.db. This
script only reads data/source/ — it never writes to it.

Validation happens entirely in memory before any database write: the whole
workbook is parsed and checked first, and only if that succeeds does the
script open a transaction and load it. A malformed workbook therefore fails
loudly and leaves any existing database untouched.

Re-running this script is safe: each run wipes and reloads the three data
tables (accounts, orders, tickets) and source_provenance inside one
transaction, so row counts never double. dataset_metadata is a single row
that gets replaced. ingestion_runs is an append-only audit log — each run
adds one row to it by design.

This script produces facts, not policy. It does not compute SLA deadlines,
cancellation fees, or service credits; see docs/architecture.md.

Usage:
    python scripts/ingest_dataset.py [--workbook PATH] [--db PATH] [--json]

Exit codes:
    0  ingestion succeeded
    1  the workbook failed validation, or ingestion otherwise failed
"""

from __future__ import annotations

import argparse
import json
import sqlite3
import sys
from dataclasses import dataclass, field
from datetime import datetime, timezone
from pathlib import Path
from typing import Any
from zoneinfo import ZoneInfo, ZoneInfoNotFoundError

import openpyxl

REPO_ROOT = Path(__file__).resolve().parent.parent
if str(REPO_ROOT) not in sys.path:
    sys.path.insert(0, str(REPO_ROOT))

from app.backend.services.database import (  # noqa: E402
    DEFAULT_DB_PATH,
    get_connection,
    initialize_schema,
)
from scripts.inspect_sources import PDF_FILES  # noqa: E402
from scripts.verify_source_pack import sha256_of  # noqa: E402

DEFAULT_WORKBOOK_PATH = REPO_ROOT / "data" / "source" / "ParcelPilot_Assessment_Data.xlsx"

INGESTION_SCRIPT_VERSION = "1.0.0"

REQUIRED_SHEETS = ("README", "accounts", "orders", "tickets")

# Exact column sets expected per sheet (order-independent; extra or missing
# columns both fail ingestion — the workbook shape is load-bearing).
ACCOUNTS_COLUMNS = (
    "account_id",
    "account_name",
    "plan",
    "status",
    "csm",
    "contract_file",
    "premium_support",
    "notes",
)
ORDERS_COLUMNS = (
    "order_id",
    "account_id",
    "carrier",
    "status",
    "booked_at",
    "pickup_window_start",
    "pickup_window_end",
    "pickup_actual_at",
    "shipment_fee_inr",
    "carrier_fault",
    "customer_fault",
    "cancellation_requested_at",
    "notes",
)
TICKETS_COLUMNS = (
    "ticket_id",
    "account_id",
    "created_at",
    "status",
    "subject",
    "description",
    "channel",
    "assigned_to",
    "last_customer_message_at",
    "historical_resolution",
)

# Columns holding row-level timestamps, parsed with the single dataset-wide
# timezone read from the workbook's own README sheet (see
# _parse_dataset_snapshot). Nothing here hard-codes a zone.
ORDERS_TIMESTAMP_COLUMNS = (
    "booked_at",
    "pickup_window_start",
    "pickup_window_end",
    "pickup_actual_at",
    "cancellation_requested_at",
)
TICKETS_TIMESTAMP_COLUMNS = ("created_at", "last_customer_message_at")

ROW_TIMESTAMP_FORMAT = "%Y-%m-%d %H:%M"

KNOWN_CONTRACT_FILES = frozenset(PDF_FILES)


class IngestionError(Exception):
    """Raised for any malformed or inconsistent source data. Never caught
    silently — a non-zero exit and a clear message beats a wrong load."""


@dataclass
class TableRecord:
    row_number: int
    values: dict[str, Any]
    # Snapshot of `values` taken immediately after reading from the
    # worksheet, before any type-coercion/timestamp-parsing mutates
    # `values` in place. This is what provenance reports — the exact cell
    # values as originally read, never the parsed/reformatted versions.
    raw_values: dict[str, Any] = field(default_factory=dict)


@dataclass
class Dataset:
    metadata: dict[str, Any]
    accounts: list[TableRecord] = field(default_factory=list)
    orders: list[TableRecord] = field(default_factory=list)
    tickets: list[TableRecord] = field(default_factory=list)


# --- workbook reading --------------------------------------------------------


def _is_blank(value: Any) -> bool:
    return value is None or value == ""


def _validate_sheets(workbook) -> None:
    actual = set(workbook.sheetnames)
    missing = [s for s in REQUIRED_SHEETS if s not in actual]
    if missing:
        raise IngestionError(f"workbook is missing required sheet(s): {', '.join(missing)}")
    unexpected = [s for s in workbook.sheetnames if s not in REQUIRED_SHEETS]
    if unexpected:
        raise IngestionError(f"workbook has unexpected sheet(s): {', '.join(unexpected)}")


def _read_sheet_header(ws, sheet_name: str, expected_columns: tuple[str, ...]) -> list[str]:
    rows = ws.iter_rows(min_row=1, max_row=1, values_only=True)
    header_row = next(rows, None)
    if header_row is None:
        raise IngestionError(f"sheet '{sheet_name}' has no header row")
    header = [str(c) if c is not None else "" for c in header_row]

    expected_set = set(expected_columns)
    actual_set = set(header)
    missing = expected_set - actual_set
    extra = actual_set - expected_set
    if missing:
        raise IngestionError(f"sheet '{sheet_name}' is missing required column(s): {', '.join(sorted(missing))}")
    if extra:
        raise IngestionError(f"sheet '{sheet_name}' has unexpected column(s): {', '.join(sorted(extra))}")
    return header


def _read_sheet_records(ws, sheet_name: str, expected_columns: tuple[str, ...]) -> list[TableRecord]:
    header = _read_sheet_header(ws, sheet_name, expected_columns)
    records: list[TableRecord] = []
    for row_number, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        if row is None or all(_is_blank(v) for v in row):
            continue  # skip fully blank trailing rows
        values = {header[i]: (row[i] if i < len(row) else None) for i in range(len(header))}
        values = {k: (None if _is_blank(v) else v) for k, v in values.items()}
        records.append(TableRecord(row_number=row_number, values=values, raw_values=dict(values)))
    return records


def _read_readme_sheet(ws) -> dict[str, Any]:
    kv: dict[str, Any] = {}
    for row in ws.iter_rows(values_only=True):
        if row and not _is_blank(row[0]) and len(row) > 1 and not _is_blank(row[1]):
            kv[str(row[0])] = row[1]
    return kv


# --- timestamp parsing --------------------------------------------------------


def _parse_dataset_snapshot(raw: str) -> tuple[datetime, str]:
    """Parse "<naive timestamp> <IANA zone>", e.g. "2026-08-16 11:00 Asia/Kolkata".

    Returns (timezone-aware datetime, zone name). This is the single source
    of the timezone applied to every row-level timestamp in the workbook.
    """
    if not isinstance(raw, str):
        raise IngestionError(f"dataset snapshot value is not text: {raw!r}")
    parts = raw.strip().rsplit(" ", 1)
    if len(parts) != 2:
        raise IngestionError(
            f"dataset snapshot does not match '<YYYY-MM-DD HH:MM> <IANA zone>': {raw!r}"
        )
    naive_str, tz_name = parts
    try:
        naive = datetime.strptime(naive_str, ROW_TIMESTAMP_FORMAT)
    except ValueError as exc:
        raise IngestionError(f"dataset snapshot timestamp is malformed: {raw!r} ({exc})") from exc
    try:
        tz = ZoneInfo(tz_name)
    except ZoneInfoNotFoundError as exc:
        raise IngestionError(f"dataset snapshot names an unknown timezone: {tz_name!r}") from exc
    return naive.replace(tzinfo=tz), tz_name


def _parse_row_timestamp(raw: Any, *, tz: ZoneInfo, sheet: str, row_number: int, column: str) -> datetime:
    if not isinstance(raw, str):
        raise IngestionError(
            f"{sheet}!row {row_number} column '{column}': expected a text timestamp, got {raw!r}"
        )
    try:
        naive = datetime.strptime(raw.strip(), ROW_TIMESTAMP_FORMAT)
    except ValueError as exc:
        raise IngestionError(
            f"{sheet}!row {row_number} column '{column}': malformed timestamp {raw!r} "
            f"(expected format 'YYYY-MM-DD HH:MM'): {exc}"
        ) from exc
    return naive.replace(tzinfo=tz)


# --- validation + typed row building -----------------------------------------


def _require_type(value: Any, expected_type: type, *, sheet: str, row_number: int, column: str) -> Any:
    if value is None:
        return None
    if expected_type is float and isinstance(value, (int, float)) and not isinstance(value, bool):
        return float(value)
    if expected_type is bool and isinstance(value, bool):
        return value
    if not isinstance(value, expected_type):
        raise IngestionError(
            f"{sheet}!row {row_number} column '{column}': expected {expected_type.__name__}, "
            f"got {type(value).__name__} ({value!r})"
        )
    return value


def _build_accounts(records: list[TableRecord]) -> list[TableRecord]:
    seen_ids: set[str] = set()
    for rec in records:
        v = rec.values
        account_id = _require_type(v["account_id"], str, sheet="accounts", row_number=rec.row_number, column="account_id")
        if not account_id:
            raise IngestionError(f"accounts!row {rec.row_number}: account_id is required")
        if account_id in seen_ids:
            raise IngestionError(f"accounts!row {rec.row_number}: duplicate account_id {account_id!r}")
        seen_ids.add(account_id)

        contract_file = _require_type(v["contract_file"], str, sheet="accounts", row_number=rec.row_number, column="contract_file")
        if contract_file is not None and contract_file not in KNOWN_CONTRACT_FILES:
            raise IngestionError(
                f"accounts!row {rec.row_number}: contract_file {contract_file!r} is not one of the "
                f"supplied source PDFs"
            )

        v["account_id"] = account_id
        v["account_name"] = _require_type(v["account_name"], str, sheet="accounts", row_number=rec.row_number, column="account_name")
        v["plan"] = _require_type(v["plan"], str, sheet="accounts", row_number=rec.row_number, column="plan")
        v["status"] = _require_type(v["status"], str, sheet="accounts", row_number=rec.row_number, column="status")
        v["csm"] = _require_type(v["csm"], str, sheet="accounts", row_number=rec.row_number, column="csm")
        v["contract_file"] = contract_file
        v["premium_support"] = _require_type(v["premium_support"], bool, sheet="accounts", row_number=rec.row_number, column="premium_support")
        v["notes"] = _require_type(v["notes"], str, sheet="accounts", row_number=rec.row_number, column="notes")
    return records


def _build_orders(records: list[TableRecord], *, known_account_ids: set[str], tz: ZoneInfo) -> list[TableRecord]:
    seen_ids: set[str] = set()
    for rec in records:
        v = rec.values
        order_id = _require_type(v["order_id"], str, sheet="orders", row_number=rec.row_number, column="order_id")
        if not order_id:
            raise IngestionError(f"orders!row {rec.row_number}: order_id is required")
        if order_id in seen_ids:
            raise IngestionError(f"orders!row {rec.row_number}: duplicate order_id {order_id!r}")
        seen_ids.add(order_id)

        account_id = _require_type(v["account_id"], str, sheet="orders", row_number=rec.row_number, column="account_id")
        if not account_id:
            raise IngestionError(f"orders!row {rec.row_number}: account_id is required")
        if account_id not in known_account_ids:
            raise IngestionError(
                f"orders!row {rec.row_number}: account_id {account_id!r} does not match any row in "
                f"the accounts sheet"
            )

        v["order_id"] = order_id
        v["account_id"] = account_id
        v["carrier"] = _require_type(v["carrier"], str, sheet="orders", row_number=rec.row_number, column="carrier")
        v["status"] = _require_type(v["status"], str, sheet="orders", row_number=rec.row_number, column="status")
        v["shipment_fee_inr"] = _require_type(v["shipment_fee_inr"], float, sheet="orders", row_number=rec.row_number, column="shipment_fee_inr")
        v["carrier_fault"] = _require_type(v["carrier_fault"], bool, sheet="orders", row_number=rec.row_number, column="carrier_fault")
        v["customer_fault"] = _require_type(v["customer_fault"], bool, sheet="orders", row_number=rec.row_number, column="customer_fault")
        v["notes"] = _require_type(v["notes"], str, sheet="orders", row_number=rec.row_number, column="notes")

        for column in ORDERS_TIMESTAMP_COLUMNS:
            raw = v[column]
            v[column] = (
                None
                if raw is None
                else _parse_row_timestamp(raw, tz=tz, sheet="orders", row_number=rec.row_number, column=column)
            )
    return records


def _build_tickets(records: list[TableRecord], *, known_account_ids: set[str], tz: ZoneInfo) -> list[TableRecord]:
    seen_ids: set[str] = set()
    for rec in records:
        v = rec.values
        ticket_id = _require_type(v["ticket_id"], str, sheet="tickets", row_number=rec.row_number, column="ticket_id")
        if not ticket_id:
            raise IngestionError(f"tickets!row {rec.row_number}: ticket_id is required")
        if ticket_id in seen_ids:
            raise IngestionError(f"tickets!row {rec.row_number}: duplicate ticket_id {ticket_id!r}")
        seen_ids.add(ticket_id)

        account_id = _require_type(v["account_id"], str, sheet="tickets", row_number=rec.row_number, column="account_id")
        if not account_id:
            raise IngestionError(f"tickets!row {rec.row_number}: account_id is required")
        if account_id not in known_account_ids:
            raise IngestionError(
                f"tickets!row {rec.row_number}: account_id {account_id!r} does not match any row in "
                f"the accounts sheet"
            )

        v["ticket_id"] = ticket_id
        v["account_id"] = account_id
        v["status"] = _require_type(v["status"], str, sheet="tickets", row_number=rec.row_number, column="status")
        v["subject"] = _require_type(v["subject"], str, sheet="tickets", row_number=rec.row_number, column="subject")
        v["description"] = _require_type(v["description"], str, sheet="tickets", row_number=rec.row_number, column="description")
        v["channel"] = _require_type(v["channel"], str, sheet="tickets", row_number=rec.row_number, column="channel")
        v["assigned_to"] = _require_type(v["assigned_to"], str, sheet="tickets", row_number=rec.row_number, column="assigned_to")
        v["historical_resolution"] = _require_type(v["historical_resolution"], str, sheet="tickets", row_number=rec.row_number, column="historical_resolution")

        for column in TICKETS_TIMESTAMP_COLUMNS:
            raw = v[column]
            v[column] = (
                None
                if raw is None
                else _parse_row_timestamp(raw, tz=tz, sheet="tickets", row_number=rec.row_number, column=column)
            )
    return records


def build_dataset(workbook_path: Path) -> Dataset:
    """Read and fully validate the workbook. Raises IngestionError on any
    problem. Performs no I/O against the database."""
    if not workbook_path.is_file():
        raise IngestionError(f"workbook not found: {workbook_path}")

    try:
        wb = openpyxl.load_workbook(workbook_path, data_only=True, read_only=True)
    except Exception as exc:
        raise IngestionError(f"workbook could not be opened: {exc}") from exc

    try:
        _validate_sheets(wb)

        readme_kv = _read_readme_sheet(wb["README"])
        snapshot_raw = readme_kv.get("Dataset snapshot")
        if snapshot_raw is None:
            raise IngestionError("README sheet is missing a 'Dataset snapshot' entry")
        snapshot_at, tz_name = _parse_dataset_snapshot(str(snapshot_raw))
        tz = ZoneInfo(tz_name)

        account_records = _build_accounts(_read_sheet_records(wb["accounts"], "accounts", ACCOUNTS_COLUMNS))
        known_account_ids = {r.values["account_id"] for r in account_records}

        order_records = _build_orders(
            _read_sheet_records(wb["orders"], "orders", ORDERS_COLUMNS),
            known_account_ids=known_account_ids,
            tz=tz,
        )
        ticket_records = _build_tickets(
            _read_sheet_records(wb["tickets"], "tickets", TICKETS_COLUMNS),
            known_account_ids=known_account_ids,
            tz=tz,
        )

        metadata = {
            "dataset_snapshot_raw": str(snapshot_raw),
            "dataset_snapshot_at": snapshot_at,
            "dataset_timezone": tz_name,
            "currency": readme_kv.get("Currency"),
            "notes": readme_kv.get("Notes"),
            "important_note": readme_kv.get("Important"),
            "source_sheet_names": list(wb.sheetnames),
        }

        return Dataset(metadata=metadata, accounts=account_records, orders=order_records, tickets=ticket_records)
    finally:
        wb.close()


# --- database loading ---------------------------------------------------------


def _iso(dt: datetime | None) -> str | None:
    return None if dt is None else dt.isoformat()


def load_into_db(
    conn: sqlite3.Connection,
    dataset: Dataset,
    *,
    workbook_path: Path,
    workbook_sha256: str,
) -> dict[str, int]:
    """Wipe and reload the data tables inside a single transaction.

    ingestion_runs is append-only; every other data table reflects exactly
    the dataset just built, nothing accumulated from prior runs.
    """
    now_utc = datetime.now(timezone.utc).isoformat()

    with conn:
        cur = conn.execute(
            """
            INSERT INTO ingestion_runs
                (started_at_utc, source_workbook_path, source_workbook_sha256,
                 ingestion_script_version, status)
            VALUES (?, ?, ?, ?, 'running')
            """,
            (now_utc, str(workbook_path), workbook_sha256, INGESTION_SCRIPT_VERSION),
        )
        run_id = cur.lastrowid

        conn.execute("DELETE FROM orders")
        conn.execute("DELETE FROM tickets")
        conn.execute("DELETE FROM source_provenance")
        conn.execute("DELETE FROM accounts")

        for rec in dataset.accounts:
            v = rec.values
            conn.execute(
                """
                INSERT INTO accounts
                    (account_id, account_name, plan, status, csm, contract_file,
                     premium_support, notes)
                VALUES (?, ?, ?, ?, ?, ?, ?, ?)
                """,
                (
                    v["account_id"], v["account_name"], v["plan"], v["status"], v["csm"],
                    v["contract_file"], v["premium_support"], v["notes"],
                ),
            )
            _insert_provenance(conn, run_id, "accounts", v["account_id"], workbook_path.name, "accounts", rec)

        for rec in dataset.orders:
            v = rec.values
            conn.execute(
                """
                INSERT INTO orders
                    (order_id, account_id, carrier, status, booked_at,
                     pickup_window_start, pickup_window_end, pickup_actual_at,
                     shipment_fee_inr, carrier_fault, customer_fault,
                     cancellation_requested_at, notes)
                VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                """,
                (
                    v["order_id"], v["account_id"], v["carrier"], v["status"],
                    _iso(v["booked_at"]), _iso(v["pickup_window_start"]), _iso(v["pickup_window_end"]),
                    _iso(v["pickup_actual_at"]), v["shipment_fee_inr"], v["carrier_fault"], v["customer_fault"],
                    _iso(v["cancellation_requested_at"]), v["notes"],
                ),
            )
            _insert_provenance(conn, run_id, "orders", v["order_id"], workbook_path.name, "orders", rec)

        for rec in dataset.tickets:
            v = rec.values
            conn.execute(
                """
                INSERT INTO tickets
                    (ticket_id, account_id, created_at, status, subject, description,
                     channel, assigned_to, last_customer_message_at, historical_resolution)
                VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                """,
                (
                    v["ticket_id"], v["account_id"], _iso(v["created_at"]), v["status"], v["subject"],
                    v["description"], v["channel"], v["assigned_to"],
                    _iso(v["last_customer_message_at"]), v["historical_resolution"],
                ),
            )
            _insert_provenance(conn, run_id, "tickets", v["ticket_id"], workbook_path.name, "tickets", rec)

        m = dataset.metadata
        conn.execute(
            """
            INSERT INTO dataset_metadata
                (id, dataset_snapshot_raw, dataset_snapshot_at, dataset_timezone,
                 currency, notes, important_note, source_workbook_filename,
                 source_workbook_sha256, source_sheet_names, ingested_at_utc,
                 ingestion_script_version)
            VALUES (1, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
            ON CONFLICT (id) DO UPDATE SET
                dataset_snapshot_raw = excluded.dataset_snapshot_raw,
                dataset_snapshot_at = excluded.dataset_snapshot_at,
                dataset_timezone = excluded.dataset_timezone,
                currency = excluded.currency,
                notes = excluded.notes,
                important_note = excluded.important_note,
                source_workbook_filename = excluded.source_workbook_filename,
                source_workbook_sha256 = excluded.source_workbook_sha256,
                source_sheet_names = excluded.source_sheet_names,
                ingested_at_utc = excluded.ingested_at_utc,
                ingestion_script_version = excluded.ingestion_script_version
            """,
            (
                m["dataset_snapshot_raw"], _iso(m["dataset_snapshot_at"]), m["dataset_timezone"],
                m["currency"], m["notes"], m["important_note"], workbook_path.name, workbook_sha256,
                json.dumps(m["source_sheet_names"]), now_utc, INGESTION_SCRIPT_VERSION,
            ),
        )

        conn.execute(
            "UPDATE ingestion_runs SET finished_at_utc = ?, status = 'success' WHERE id = ?",
            (datetime.now(timezone.utc).isoformat(), run_id),
        )

    return {
        "accounts": len(dataset.accounts),
        "orders": len(dataset.orders),
        "tickets": len(dataset.tickets),
    }


def _insert_provenance(
    conn: sqlite3.Connection,
    run_id: int,
    target_table: str,
    target_id: str,
    source_file: str,
    source_sheet: str,
    rec: TableRecord,
) -> None:
    raw_row = {k: (v.isoformat() if isinstance(v, datetime) else v) for k, v in rec.raw_values.items()}
    conn.execute(
        """
        INSERT INTO source_provenance
            (ingestion_run_id, target_table, target_id, source_file, source_sheet,
             source_row_number, raw_row_json)
        VALUES (?, ?, ?, ?, ?, ?, ?)
        """,
        (run_id, target_table, target_id, source_file, source_sheet, rec.row_number, json.dumps(raw_row, sort_keys=True)),
    )


# --- orchestration -------------------------------------------------------------


def ingest(
    workbook_path: Path = DEFAULT_WORKBOOK_PATH,
    db_path: Path = DEFAULT_DB_PATH,
) -> dict[str, Any]:
    workbook_path = Path(workbook_path)
    db_path = Path(db_path)

    dataset = build_dataset(workbook_path)
    workbook_sha256 = sha256_of(workbook_path)

    conn = get_connection(db_path)
    try:
        initialize_schema(conn)
        row_counts = load_into_db(conn, dataset, workbook_path=workbook_path, workbook_sha256=workbook_sha256)
    finally:
        conn.close()

    return {
        "ok": True,
        "workbook": str(workbook_path),
        "workbook_sha256": workbook_sha256,
        "database": str(db_path),
        "row_counts": row_counts,
        "dataset_snapshot": dataset.metadata["dataset_snapshot_raw"],
    }


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--workbook", type=Path, default=DEFAULT_WORKBOOK_PATH)
    parser.add_argument("--db", type=Path, default=DEFAULT_DB_PATH)
    parser.add_argument("--json", action="store_true")
    args = parser.parse_args(argv)

    try:
        result = ingest(workbook_path=args.workbook, db_path=args.db)
    except IngestionError as exc:
        error = {"ok": False, "error": str(exc)}
        if args.json:
            print(json.dumps(error, indent=2))
        else:
            print(f"INGESTION FAILED: {exc}")
        return 1

    if args.json:
        print(json.dumps(result, indent=2, default=str))
    else:
        print(f"Workbook: {result['workbook']}")
        print(f"  sha256: {result['workbook_sha256']}")
        print(f"Database: {result['database']}")
        print(f"Dataset snapshot: {result['dataset_snapshot']}")
        for table, count in result["row_counts"].items():
            print(f"  {table}: {count} row(s)")
        print("INGESTION OK")
    return 0


if __name__ == "__main__":
    sys.exit(main())
